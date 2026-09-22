"""Small, source-preserving parsers for PDF, DOCX and XLSX documents."""

from __future__ import annotations

import hashlib
import re
from pathlib import Path
from typing import Any
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from pydantic import BaseModel, ConfigDict, Field, model_validator

try:
    from .domain import SourcePointer
except ImportError:
    from domain import SourcePointer  # type: ignore


class IngestionError(RuntimeError):
    """The source cannot be parsed defensibly."""


MAX_DOCUMENT_BYTES = 25 * 1024 * 1024
MAX_XML_BYTES = 10 * 1024 * 1024


class DocumentChunk(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    chunk_id: str = Field(min_length=1)
    document_id: str = Field(min_length=1)
    document_name: str = Field(min_length=1)
    document_type: str = Field(min_length=1)
    text: str = Field(min_length=1)
    content_hash: str = Field(pattern=r"^sha256:[0-9a-f]{64}$")
    source_pointer: SourcePointer
    vendor_id: str | None = None
    proposal_id: str | None = None

    @model_validator(mode="after")
    def enforce_scope_pair(self) -> "DocumentChunk":
        if (self.vendor_id is None) != (self.proposal_id is None):
            raise ValueError("vendor_id and proposal_id must be supplied together")
        return self


def _hash(text: str) -> str:
    return f"sha256:{hashlib.sha256(text.encode('utf-8')).hexdigest()}"


def _chunk(document_id: str, name: str, kind: str, text: str, pointer_data: dict[str, Any], vendor_id: str | None, proposal_id: str | None, index: int) -> DocumentChunk:
    clean = re.sub(r"\s+", " ", text).strip()
    if not clean:
        raise IngestionError(f"empty source chunk at index {index}")
    content_hash = _hash(clean)
    pointer = SourcePointer(source_pointer_id=f"PTR_{document_id}_{index:04d}", document_id=document_id, document_name=name, document_type=kind, content_hash=content_hash, resolvable=True, **pointer_data)
    return DocumentChunk(chunk_id=f"CHK_{document_id}_{index:04d}", document_id=document_id, document_name=name, document_type=kind, text=clean, content_hash=content_hash, source_pointer=pointer, vendor_id=vendor_id, proposal_id=proposal_id)


def parse_pdf(path: Path, document_id: str, vendor_id: str | None = None, proposal_id: str | None = None) -> list[DocumentChunk]:
    try:
        from pypdf import PdfReader
        pages = PdfReader(str(path)).pages
    except Exception as exc:
        raise IngestionError(f"PDF parsing failed: {path.name}") from exc
    chunks: list[DocumentChunk] = []
    for page_number, page in enumerate(pages, start=1):
        try:
            text = page.extract_text() or ""
        except Exception as exc:
            raise IngestionError(f"PDF page extraction failed: {page_number}") from exc
        if text.strip():
            chunks.append(_chunk(document_id, path.name, "PDF", text, {"page_number": page_number}, vendor_id, proposal_id, len(chunks)))
    if not chunks:
        raise IngestionError("PDF contains no extractable text")
    return chunks


def parse_docx(path: Path, document_id: str, vendor_id: str | None = None, proposal_id: str | None = None) -> list[DocumentChunk]:
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    try:
        if path.stat().st_size > MAX_DOCUMENT_BYTES:
            raise IngestionError("DOCX exceeds the 25 MiB MVP limit")
        with ZipFile(path) as archive:
            document_member = archive.getinfo("word/document.xml")
            if document_member.file_size > MAX_XML_BYTES:
                raise IngestionError("DOCX XML part exceeds the parser limit")
            root = ET.fromstring(archive.read("word/document.xml"))
    except Exception as exc:
        raise IngestionError(f"DOCX parsing failed: {path.name}") from exc
    chunks: list[DocumentChunk] = []
    for paragraph_number, paragraph in enumerate(root.findall(".//w:p", ns), start=1):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", ns))
        if text.strip():
            chunks.append(_chunk(document_id, path.name, "DOCX", text, {"section": f"paragraph {paragraph_number}"}, vendor_id, proposal_id, len(chunks)))
    if not chunks:
        raise IngestionError("DOCX contains no extractable text")
    return chunks


def parse_xlsx(path: Path, document_id: str, vendor_id: str | None = None, proposal_id: str | None = None) -> list[DocumentChunk]:
    try:
        if path.stat().st_size > MAX_DOCUMENT_BYTES:
            raise IngestionError("XLSX exceeds the 25 MiB MVP limit")
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise IngestionError(f"XLSX parsing failed: {path.name}") from exc
    chunks: list[DocumentChunk] = []
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                chunks.append(_chunk(document_id, path.name, "XLSX", " | ".join(values), {"sheet_name": sheet.title, "row_start": row_number, "row_end": row_number}, vendor_id, proposal_id, len(chunks)))
    if not chunks:
        raise IngestionError("XLSX contains no non-empty cells")
    return chunks


def parse_document(path: str | Path, document_id: str, media_type: str, vendor_id: str | None = None, proposal_id: str | None = None) -> list[DocumentChunk]:
    file_path = Path(path)
    if file_path.exists() and file_path.stat().st_size > MAX_DOCUMENT_BYTES:
        raise IngestionError("document exceeds the 25 MiB MVP limit")
    if (vendor_id is None) != (proposal_id is None):
        raise ValueError("vendor_id and proposal_id must be supplied together")
    if media_type == "application/pdf":
        return parse_pdf(file_path, document_id, vendor_id, proposal_id)
    if media_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return parse_docx(file_path, document_id, vendor_id, proposal_id)
    if media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return parse_xlsx(file_path, document_id, vendor_id, proposal_id)
    raise IngestionError(f"unsupported media type: {media_type}")
